"""Offline Phase 2 performance and Python-allocation baseline harness.

Absolute timings are machine-specific observations, not a public performance SLA.
Workbook generation is deterministic and occurs outside every measured region.
"""

from __future__ import annotations

import argparse
import gc
import json
import platform
from pathlib import Path
import statistics
import tempfile
import time
import tracemalloc
from typing import Callable, TypeVar

from openpyxl import Workbook

from xlsform_lint.diagnostics import Diagnostic, sort_diagnostics
from xlsform_lint.engine import lint_workbook
from xlsform_lint.config import ConfigSettings
from xlsform_lint.json_formatter import format_json
from xlsform_lint.native_model import build_native_model
from xlsform_lint.native_rules import run_native_rules
from xlsform_lint.policy import CliPolicyOptions, apply_policy, resolve_policy
from xlsform_lint.pyxform_adapter import validate_workbook
from xlsform_lint.sarif_formatter import format_sarif
from xlsform_lint.text_formatter import format_diagnostics
from xlsform_lint.workbook_input import read_workbook_bytes


ROOT = Path(__file__).parents[1]
CORPUS = ROOT / "tests" / "fixtures" / "corpus"
T = TypeVar("T")
SERIES = ((50, 100), (150, 300), (300, 600), (500, 1000))
QUICK_SERIES = ((20, 40), (40, 80), (60, 120))


def write_synthetic_workbook(path: Path, survey_rows: int, choice_rows: int) -> Path:
    """Write a clean, deterministic workload with mixed representative constructs."""

    if survey_rows < 3 or choice_rows < 4:
        raise ValueError("synthetic workload needs at least 3 survey and 4 choice rows")
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append(
        (
            "type",
            "name",
            "label::English",
            "label::French",
            "constraint",
            "constraint_message::English",
            "constraint_message::French",
            "relevant",
            "calculation",
            "choice_filter",
        )
    )
    survey.append(("select_one districts", "district", "District", "District FR"))
    survey.append(("begin group", "details", "Details", "Détails"))
    for index in range(survey_rows - 3):
        name = f"question_{index:04d}"
        kind = index % 5
        if kind == 0:
            row = ("select_one options", name, f"Option {index}", f"Option FR {index}", None, None, None, None, None, "district=${district}")
        elif kind == 1:
            row = ("integer", name, f"Number {index}", f"Nombre {index}", ". >= 0", "Enter zero or more", "Entrez zéro ou plus")
        elif kind == 2:
            row = ("text", name, f"Detail {index}", f"Détail {index}", None, None, None, "${district} != ''")
        elif kind == 3:
            row = ("calculate", name, None, None, None, None, None, None, "1 + 1")
        else:
            row = ("note", name, f"Note {index}", f"Note FR {index}")
        survey.append(row)
    survey.append(("end group",))

    choices = workbook.create_sheet("choices")
    choices.append(("list_name", "name", "label::English", "label::French", "district"))
    choices.append(("districts", "north", "North", "Nord", None))
    choices.append(("districts", "south", "South", "Sud", None))
    for index in range(choice_rows - 2):
        choices.append(
            (
                "options",
                f"option_{index:05d}",
                f"Option {index}",
                f"Option FR {index}",
                "north" if index % 2 == 0 else "south",
            )
        )

    settings = workbook.create_sheet("settings")
    settings.append(("form_title", "form_id", "version"))
    settings.append(("Synthetic performance form", "synthetic_performance", "1"))
    workbook.save(path)
    workbook.close()
    return path


def _measure(
    operation: Callable[[], T], *, warmups: int, iterations: int
) -> tuple[dict[str, float], T]:
    for _ in range(warmups):
        operation()
    samples: list[float] = []
    results: list[T] = []
    for _ in range(iterations):
        gc.collect()
        started = time.perf_counter()
        results.append(operation())
        samples.append(time.perf_counter() - started)
    if any(result != results[0] for result in results[1:]):
        raise RuntimeError("benchmark operation produced nondeterministic results")
    return {
        "median_seconds": statistics.median(samples),
        "min_seconds": min(samples),
        "max_seconds": max(samples),
    }, results[0]


def _peak_python_bytes(operation: Callable[[], object]) -> int:
    gc.collect()
    tracemalloc.start()
    try:
        operation()
        _, peak = tracemalloc.get_traced_memory()
        return peak
    finally:
        tracemalloc.stop()


def measure_workload(
    name: str, path: Path, *, survey_rows: int, choice_rows: int, warmups: int, iterations: int
) -> dict[str, object]:
    """Measure existing boundaries without adding runtime instrumentation."""

    workbook_bytes = read_workbook_bytes(path)
    model = build_native_model(workbook_bytes)
    native_diagnostics = run_native_rules(model, str(path))
    pyxform_diagnostic = validate_workbook(workbook_bytes, str(path))
    diagnostics = sort_diagnostics(
        native_diagnostics + (() if pyxform_diagnostic is None else (pyxform_diagnostic,))
    )
    policy = resolve_policy(ConfigSettings(), CliPolicyOptions())

    stages: dict[str, dict[str, float]] = {}
    stages["byte_acquisition"], measured_bytes = _measure(
        lambda: read_workbook_bytes(path), warmups=warmups, iterations=iterations
    )
    if measured_bytes != workbook_bytes:
        raise RuntimeError("workbook bytes changed during measurement")
    stages["source_model"], measured_model = _measure(
        lambda: build_native_model(workbook_bytes), warmups=warmups, iterations=iterations
    )
    stages["native_rules"], measured_native = _measure(
        lambda: run_native_rules(model, str(path)), warmups=warmups, iterations=iterations
    )
    stages["pyxform"], measured_px = _measure(
        lambda: validate_workbook(workbook_bytes, str(path)), warmups=warmups, iterations=iterations
    )
    stages["combine_sort"], measured_combined = _measure(
        lambda: sort_diagnostics(
            native_diagnostics + (() if pyxform_diagnostic is None else (pyxform_diagnostic,))
        ),
        warmups=warmups,
        iterations=iterations,
    )
    stages["policy"], measured_policy = _measure(
        lambda: apply_policy(diagnostics, policy), warmups=warmups, iterations=iterations
    )
    stages["text_format"], _ = _measure(
        lambda: format_diagnostics(diagnostics), warmups=warmups, iterations=iterations
    )
    stages["json_format"], _ = _measure(
        lambda: format_json(str(path), diagnostics), warmups=warmups, iterations=iterations
    )
    stages["sarif_format"], _ = _measure(
        lambda: format_sarif(str(path), diagnostics), warmups=warmups, iterations=iterations
    )
    overall, measured_diagnostics = _measure(
        lambda: lint_workbook(path), warmups=warmups, iterations=iterations
    )
    if (
        len(measured_model.survey_rows) != survey_rows
        or len(measured_model.choice_rows) != choice_rows
        or measured_native != native_diagnostics
        or measured_px != pyxform_diagnostic
        or measured_combined != diagnostics
        or measured_policy != diagnostics
        or measured_diagnostics != diagnostics
    ):
        raise RuntimeError("stage and overall correctness results disagree")
    return {
        "name": name,
        "survey_rows": survey_rows,
        "choice_rows": choice_rows,
        "xlsx_bytes": len(workbook_bytes),
        "diagnostic_count": len(diagnostics),
        "rule_ids": [item.rule_id for item in diagnostics],
        "overall": overall,
        "stages": stages,
        "peak_python_traced_bytes": _peak_python_bytes(lambda: lint_workbook(path)),
    }


def run_benchmark(*, warmups: int = 1, iterations: int = 5, quick: bool = False) -> dict[str, object]:
    """Generate workloads, measure them, and return a machine-readable report."""

    if warmups < 0 or iterations < 1:
        raise ValueError("warmups must be nonnegative and iterations must be positive")
    series = QUICK_SERIES if quick else SERIES
    with tempfile.TemporaryDirectory(prefix="xlsform-lint-phase2h-") as directory:
        generated_root = Path(directory)
        generated = []
        for survey_rows, choice_rows in series:
            path = write_synthetic_workbook(
                generated_root / f"synthetic-{survey_rows}-{choice_rows}.xlsx",
                survey_rows,
                choice_rows,
            )
            generated.append((path, survey_rows, choice_rows))

        small_path = CORPUS / "valid" / "basic_clean.xlsx"
        medium_path = CORPUS / "valid" / "large_choice_list.xlsx"
        large_path, large_survey, large_choices = generated[-1]
        workloads = (
            ("small", small_path, 1, 0),
            ("medium", medium_path, 1, 250),
            ("large", large_path, large_survey, large_choices),
        )
        measured = [
            measure_workload(
                name,
                path,
                survey_rows=survey_rows,
                choice_rows=choice_rows,
                warmups=warmups,
                iterations=iterations,
            )
            for name, path, survey_rows, choice_rows in workloads
        ]
        scaling = [
            measure_workload(
                f"S{index}",
                path,
                survey_rows=survey_rows,
                choice_rows=choice_rows,
                warmups=0,
                iterations=1 if quick else max(2, iterations // 2),
            )
            for index, (path, survey_rows, choice_rows) in enumerate(generated, 1)
        ]
    return {
        "environment": {
            "platform": platform.platform(),
            "system": platform.system(),
            "machine": platform.machine(),
            "python": platform.python_version(),
            "processor": platform.processor() or "unavailable",
        },
        "method": {
            "warmups": warmups,
            "iterations": iterations,
            "clock": "time.perf_counter",
            "summary": "median/min/max",
            "memory": "tracemalloc peak Python-tracked bytes; not process RSS",
            "generation_timed": False,
            "stage_warning": "isolated stage medians are not additive and may double-count setup",
            "public_sla": False,
        },
        "series": [{"survey_rows": survey, "choice_rows": choices} for survey, choices in series],
        "workloads": measured,
        "scaling": scaling,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--warmups", type=int, default=1)
    parser.add_argument("--iterations", type=int, default=5)
    parser.add_argument("--quick", action="store_true", help="use reduced structural workloads")
    parser.add_argument("--output", type=Path, help="write JSON report to this path")
    arguments = parser.parse_args()
    report = run_benchmark(
        warmups=arguments.warmups, iterations=arguments.iterations, quick=arguments.quick
    )
    rendered = json.dumps(report, indent=2, sort_keys=True) + "\n"
    if arguments.output is None:
        print(rendered, end="")
    else:
        arguments.output.write_text(rendered, encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
