from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from scripts.benchmark_phase2 import run_benchmark, write_synthetic_workbook
from xlsform_lint import engine
from xlsform_lint.engine import lint_workbook


def test_generated_workload_has_requested_shape_and_is_clean_deterministic(tmp_path: Path) -> None:
    path = write_synthetic_workbook(tmp_path / "generated.xlsx", 120, 240)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        assert workbook["survey"].max_row == 121
        assert workbook["choices"].max_row == 241
        assert workbook.sheetnames == ["survey", "choices", "settings"]
    finally:
        workbook.close()
    assert lint_workbook(path) == lint_workbook(path) == ()


def test_one_lint_reads_bytes_and_invokes_pyxform_once(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    path = write_synthetic_workbook(tmp_path / "read-once.xlsx", 30, 60)
    reads = conversions = 0
    original_read = engine.read_workbook_bytes
    original_validate = engine.validate_workbook

    def counted_read(source: Path) -> bytes:
        nonlocal reads
        reads += 1
        return original_read(source)

    def counted_validate(workbook_bytes: bytes, source_path: str):
        nonlocal conversions
        conversions += 1
        return original_validate(workbook_bytes, source_path)

    monkeypatch.setattr(engine, "read_workbook_bytes", counted_read)
    monkeypatch.setattr(engine, "validate_workbook", counted_validate)
    assert lint_workbook(path) == ()
    assert reads == conversions == 1


def test_quick_harness_executes_offline_with_structural_measurements() -> None:
    report = run_benchmark(warmups=0, iterations=1, quick=True)
    assert report["method"]["public_sla"] is False
    assert report["method"]["generation_timed"] is False
    assert [item["name"] for item in report["workloads"]] == ["small", "medium", "large"]
    assert [(item["survey_rows"], item["choice_rows"]) for item in report["scaling"]] == [
        (20, 40),
        (40, 80),
        (60, 120),
    ]
    expected_stages = {
        "byte_acquisition",
        "source_model",
        "native_rules",
        "pyxform",
        "combine_sort",
        "policy",
        "text_format",
        "json_format",
        "sarif_format",
    }
    for item in (*report["workloads"], *report["scaling"]):
        assert item["diagnostic_count"] == 0 and item["rule_ids"] == []
        assert item["xlsx_bytes"] > 0 and item["peak_python_traced_bytes"] > 0
        assert set(item["stages"]) == expected_stages
        assert item["overall"]["median_seconds"] >= 0


def test_benchmark_command_emits_parseable_report_without_tracked_outputs() -> None:
    result = subprocess.run(
        [
            sys.executable,
            "scripts/benchmark_phase2.py",
            "--quick",
            "--warmups",
            "0",
            "--iterations",
            "1",
        ],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert result.returncode == 0 and result.stderr == ""
    payload = json.loads(result.stdout)
    assert payload["environment"]["python"]
    assert len(payload["workloads"]) == len(payload["scaling"]) == 3
