from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from xlsform_lint.engine import lint_workbook


CORPUS = Path(__file__).parent / "fixtures" / "corpus"
MANIFEST = json.loads((CORPUS / "manifest.json").read_text(encoding="utf-8"))
FIXTURES = MANIFEST["fixtures"]


@pytest.mark.parametrize("entry", FIXTURES, ids=lambda entry: entry["path"])
def test_manifest_fixture_exists_opens_and_has_expected_diagnostics(entry):
    path = CORPUS / entry["path"]
    assert path.is_file()
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        assert "survey" in workbook.sheetnames
        assert "settings" in workbook.sheetnames
    finally:
        workbook.close()

    actual = [diagnostic.rule_id for diagnostic in lint_workbook(path)]
    assert set(actual) == set(entry["expected_rule_ids"])
    if entry["pyxform_accept"]:
        assert not any(rule_id.startswith("PX") for rule_id in actual)
    if entry["clean"]:
        assert entry["category"] == "valid"
        assert entry["pyxform_accept"] is True
        assert actual == []


def test_manifest_is_complete_synchronized_and_categories_are_isolated():
    referenced = {entry["path"] for entry in FIXTURES}
    present = {
        path.relative_to(CORPUS).as_posix()
        for path in CORPUS.rglob("*.xlsx")
    }
    assert referenced == present
    assert len(FIXTURES) == 13
    assert {entry["category"] for entry in FIXTURES} == {"valid", "findings", "malformed"}
    assert all(entry["path"].startswith(f"{entry['category']}/") for entry in FIXTURES)
    assert all(entry["purpose"] and entry["constructs"] for entry in FIXTURES)
    assert MANIFEST["license"] == "Apache-2.0"
    assert "synthetic/minimized" in MANIFEST["provenance"]


def test_required_construct_coverage_is_declared():
    covered = {construct for entry in FIXTURES for construct in entry["constructs"]}
    required = {
        "groups", "nested_groups", "repeats", "nested_group_repeat", "select_one",
        "select_multiple", "multiple_choice_lists", "choice_filter",
        "earlier_field_choice_filter", "translations", "multiple_label_languages",
        "constraints", "constraint_messages", "relevance", "calculates", "notes",
        "hidden_internal", "required", "defaults", "large_static_choice_list",
        "mixed_legitimate", "mixed_intentional_findings",
    }
    assert required <= covered


def test_all_seven_native_rules_have_positive_corpus_oracles():
    expected = {rule for entry in FIXTURES for rule in entry["expected_rule_ids"]}
    assert {"CHO001", "LBL001", "LBL002", "I18N001", "UX001", "ORD001", "UX002"} <= expected


def test_large_choice_list_remains_bounded_and_representative():
    path = CORPUS / "valid" / "large_choice_list.xlsx"
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook["choices"].max_row == 251
    finally:
        workbook.close()
