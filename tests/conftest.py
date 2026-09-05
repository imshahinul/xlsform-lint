from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

import pytest
from openpyxl import Workbook


@pytest.fixture
def workbook_factory():
    def create(path: Path, *, question_type: str = "text") -> Path:
        workbook = Workbook()
        survey = workbook.active
        survey.title = "Survey"
        survey.append(("type", "name", "label"))
        survey.append((question_type, "person_name", "Person name"))
        settings = workbook.create_sheet("settings")
        settings.append(("form_title", "form_id", "version"))
        settings.append(("Kernel Proof", "kernel_proof", "1"))
        workbook.save(path)
        workbook.close()
        return path

    return create


@pytest.fixture
def xlsform_factory():
    """Create compact real XLSX forms with caller-controlled rows and casing."""

    def create(
        path: Path,
        survey_rows: Sequence[Sequence[object]],
        *,
        choices_rows: Sequence[Sequence[object]] | None = None,
        settings_rows: Sequence[Sequence[object]] | None = None,
        survey_title: str = "Survey",
        choices_title: str = "Choices",
    ) -> Path:
        workbook = Workbook()
        survey = workbook.active
        survey.title = survey_title
        for row in survey_rows:
            survey.append(tuple(row))
        if choices_rows is not None:
            choices = workbook.create_sheet(choices_title)
            for row in choices_rows:
                choices.append(tuple(row))
        settings = workbook.create_sheet("Settings")
        for row in settings_rows or (
            ("form_title", "form_id", "version"),
            ("Phase 2 Proof", "phase_2_proof", "1"),
        ):
            settings.append(tuple(row))
        workbook.save(path)
        workbook.close()
        return path

    return create
