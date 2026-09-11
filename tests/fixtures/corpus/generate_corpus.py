"""Generate the synthetic Phase 2 XLSForm regression corpus.

Run from the repository root with:

    python tests/fixtures/corpus/generate_corpus.py

The workbooks are intentionally small except for one 250-row static choice list.
No source form or participant data was used.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from pathlib import Path

from openpyxl import Workbook


ROOT = Path(__file__).parent
SURVEY = (
    "type",
    "name",
    "label",
    "label::English (en)",
    "label::French (fr)",
    "hint",
    "required",
    "default",
    "relevant",
    "constraint",
    "constraint_message",
    "constraint_message::English (en)",
    "constraint_message::French (fr)",
    "calculation",
    "choice_filter",
)
CHOICES = (
    "list_name",
    "name",
    "label",
    "label::English (en)",
    "label::French (fr)",
    "region",
    "area",
)


def row(
    type_: object,
    name: object,
    label: object = None,
    en: object = None,
    fr: object = None,
    hint: object = None,
    required: object = None,
    default: object = None,
    relevant: object = None,
    constraint: object = None,
    constraint_message: object = None,
    constraint_en: object = None,
    constraint_fr: object = None,
    calculation: object = None,
    choice_filter: object = None,
) -> tuple[object, ...]:
    if label is not None:
        en = label if en is None else en
        fr = label if fr is None else fr
    return (
        type_, name, label, en, fr, hint, required, default, relevant, constraint,
        constraint_message, constraint_en, constraint_fr, calculation, choice_filter,
    )


def choice(
    list_name: object,
    name: object,
    label: object = None,
    en: object = None,
    fr: object = None,
    region: object = None,
    area: object = None,
) -> tuple[object, ...]:
    if label is not None:
        en = label if en is None else en
        fr = label if fr is None else fr
    return list_name, name, label, en, fr, region, area


def write_form(
    relative_path: str,
    survey_rows: Iterable[Sequence[object]],
    choices_rows: Iterable[Sequence[object]] | None = None,
) -> None:
    path = ROOT / relative_path
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append(SURVEY)
    for survey_row in survey_rows:
        survey.append(tuple(survey_row))
    if choices_rows is not None:
        choices_sheet = workbook.create_sheet("choices")
        choices_sheet.append(CHOICES)
        for choices_row in choices_rows:
            choices_sheet.append(tuple(choices_row))
    settings = workbook.create_sheet("settings")
    settings.append(("form_title", "form_id", "version", "default_language"))
    settings.append((path.stem.replace("_", " ").title(), path.stem, "1", "English (en)"))
    workbook.save(path)
    workbook.close()


def generate() -> None:
    write_form("valid/basic_clean.xlsx", [row("text", "respondent", "Respondent")])

    write_form(
        "valid/nested_groups.xlsx",
        [
            row("begin group", "household", "Household"),
            row("text", "household_id", "Household ID", required="yes"),
            row("begin group", "location", "Location"),
            row("geopoint", "coordinates", "Coordinates"),
            row("end group", None),
            row("end group", None),
        ],
    )

    write_form(
        "valid/repeat_interaction.xlsx",
        [
            row("begin group", "roster", "Roster"),
            row("integer", "member_count", "Member count", constraint=". >= 0", constraint_message="Cannot be negative"),
            row("begin repeat", "members", "Members"),
            row("begin group", "identity", "Identity"),
            row("text", "member_name", "Member name", required="yes"),
            row("integer", "member_age", "Member age", constraint=". >= 0 and . <= 120", constraint_message="Enter an age from 0 to 120"),
            row("end group", None),
            row("end repeat", None),
            row("end group", None),
        ],
    )

    write_form(
        "valid/selects_and_filter.xlsx",
        [
            row("select_one regions", "region", "Region"),
            row("select_one facilities", "facility", "Facility", choice_filter="region=${region}"),
            row("select_multiple services", "services", "Services"),
        ],
        [
            choice("regions", "north", "North"), choice("regions", "south", "South"),
            choice("facilities", "north_clinic", "North clinic", region="north"),
            choice("facilities", "south_clinic", "South clinic", region="south"),
            choice("services", "water", "Water"), choice("services", "health", "Health"),
        ],
    )

    write_form(
        "valid/multilingual_constraints.xlsx",
        [
            row("integer", "age", en="Age", fr="Âge", required="yes", constraint=". >= 0 and . <= 120", constraint_en="Enter an age from 0 to 120", constraint_fr="Saisissez un âge de 0 à 120"),
            row("text", "nickname", en="Nickname", fr="Surnom"),
        ],
    )

    write_form(
        "valid/logic_and_internal.xlsx",
        [
            row("integer", "birth_year", "Birth year"),
            row("calculate", "approx_age", calculation="2026 - ${birth_year}"),
            row("note", "adult_note", "Adult follow-up", relevant="${approx_age} >= 18"),
            row("hidden", "case_seed", default="synthetic"),
            row("text", "adult_detail", "Adult detail", relevant="${approx_age} >= 18"),
        ],
    )

    write_form(
        "valid/required_and_defaults.xlsx",
        [
            row("date", "visit_date", "Visit date", required="yes"),
            row("text", "country_code", "Country code", default="XX"),
            row("integer", "round", "Round", default="1"),
        ],
    )

    write_form(
        "valid/large_choice_list.xlsx",
        [row("select_one products", "product", "Product")],
        [choice("products", f"product_{number:03d}", f"Synthetic product {number:03d}") for number in range(1, 251)],
    )

    write_form(
        "valid/complex_combined.xlsx",
        [
            row("select_one areas", "area", en="Area", fr="Zone", required="yes"),
            row("begin group", "visit", en="Visit", fr="Visite"),
            row("select_multiple assets", "assets", en="Assets", fr="Biens"),
            row("integer", "score", en="Score", fr="Score", constraint=". >= 0", constraint_en="Must be non-negative", constraint_fr="Doit être positif"),
            row("calculate", "double_score", calculation="${score} * 2"),
            row("note", "high_note", en="High score", fr="Score élevé", relevant="${double_score} > 10"),
            row("begin repeat", "observations", en="Observations", fr="Observations"),
            row("select_one sites", "site", en="Site", fr="Site", choice_filter="area=${area}"),
            row("text", "comment", en="Comment", fr="Commentaire", default="none"),
            row("end repeat", None),
            row("end group", None),
        ],
        [
            choice("areas", "east", en="East", fr="Est"), choice("areas", "west", en="West", fr="Ouest"),
            choice("assets", "radio", en="Radio", fr="Radio"), choice("assets", "phone", en="Phone", fr="Téléphone"),
            choice("sites", "east_1", en="East site", fr="Site est", area="east"),
            choice("sites", "west_1", en="West site", fr="Site ouest", area="west"),
        ],
    )

    write_form(
        "findings/mixed_native_findings.xlsx",
        [
            row("integer", "age", None, en="Age", constraint=". >= 0"),
            row("select_one used", "selection", "Selection", required="yes", default="a"),
            row("select_one filtered", "filtered", "Filtered", choice_filter="region=${later}"),
            row("text", "later", "Later"),
            row("note", "missing_note", hint="Intentional hint-only note"),
        ],
        [
            choice("used", "a", "A"), choice("filtered", "x", "X", region="x"),
            choice("unused", "orphan"),
        ],
    )

    write_form(
        "findings/duplicate_choice.xlsx",
        [row("select_one colors", "color", "Color")],
        [choice("colors", "red", "Red"), choice("colors", "red", "Crimson")],
    )

    write_form("malformed/unmatched_end_group.xlsx", [row("end group", None)])
    write_form(
        "malformed/unknown_choice_list.xlsx",
        [row("select_one absent", "selection", "Selection")],
        [choice("other", "value", "Value")],
    )


if __name__ == "__main__":
    generate()
