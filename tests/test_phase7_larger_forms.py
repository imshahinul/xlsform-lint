from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook

from xlsform_lint.cli import main
from xlsform_lint.engine import lint_workbook


HEADER = (
    "type", "name", "label::English", "label::French", "hint", "constraint",
    "constraint_message", "required", "default", "calculation", "choice_filter",
)


def _write_large(path: Path, *, lint: bool = False, fatal: bool = False) -> Path:
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append(HEADER)
    survey.append(("text", "district", "District", "District FR", None, None, None, None, None, None, None))
    survey.append(("select_one villages", "village", "Village", "Village FR", None, None, None, None, None, None, "district=${district}"))
    survey.append(("begin group", "household", "Household", "Ménage", None, None, None, None, None, None, None))
    survey.append(("integer", "age", "Age", "Âge", None, ". >= 0", "Enter a non-negative age", "yes", None, None, None))
    survey.append(("select_multiple amenities", "amenities", "Amenities", "Services", None, None, None, None, None, None, None))
    survey.append(("note", "guidance", "Review the answers", "Vérifier les réponses", None, None, None, None, None, None, None))
    survey.append(("calculate", "adult", None, None, None, None, None, None, None, "${age} >= 18", None))
    survey.append(("hidden", "internal_code", None, None, None, None, None, None, "v1", None, None))
    survey.append(("end group", None, None, None, None, None, None, None, None, None, None))
    survey.append(("begin repeat", "members", "Members", "Membres", None, None, None, None, None, None, None))
    survey.append(("text", "member_name", "Member name", "Nom du membre", None, None, None, None, None, None, None))
    survey.append(("end repeat", None, None, None, None, None, None, None, None, None, None))
    for number in range(1, 61):
        name = "1fatal" if fatal and number == 60 else f"detail_{number:03d}"
        label_en = None if lint and number == 1 else f"Detail {number}"
        label_fr = None if lint and number in (1, 2) else f"Détail {number}"
        hint = "Intentionally label-free native finding" if lint and number == 1 else None
        constraint = ". != ''" if lint and number == 3 else None
        message = None if constraint else None
        required = "yes" if lint and number == 4 else None
        default = "seed" if required else None
        survey.append(("text", name, label_en, label_fr, hint, constraint, message, required, default, None, None))

    choices = workbook.create_sheet("choices")
    choices.append(("list_name", "name", "label::English", "label::French", "district"))
    for number in range(1, 81):
        choices.append(("villages", f"v{number:03d}", f"Village {number}", f"Village FR {number}", "north" if number % 2 else "south"))
    choices.append(("amenities", "water", "Water", "Eau", None))
    choices.append(("amenities", "power", "Power", "Électricité", None))
    if lint:
        choices.append(("unused", "unused", "Unused", "Inutilisé", None))

    settings = workbook.create_sheet("settings")
    settings.append(("form_title", "form_id", "version"))
    settings.append(("Synthetic larger form", "synthetic_larger", "1"))
    workbook.save(path)
    workbook.close()
    return path


def test_larger_clean_and_realistic_integrated_form_is_deterministic(tmp_path, capsys) -> None:
    path = _write_large(tmp_path / "larger-clean.xlsx")
    assert lint_workbook(path) == ()
    assert main([str(path)]) == 0
    first = capsys.readouterr()
    assert main([str(path)]) == 0
    second = capsys.readouterr()
    assert first.out == second.out == "0 errors, 0 warnings\n"
    assert first.err == second.err == ""


def test_larger_lint_has_exact_findings_without_collateral_noise(tmp_path) -> None:
    diagnostics = lint_workbook(_write_large(tmp_path / "larger-lint.xlsx", lint=True))
    assert [item.rule_id for item in diagnostics] == ["LBL001", "I18N001", "UX001", "UX002", "CHO001"]
    assert Counter(item.rule_id for item in diagnostics) == Counter(
        {"CHO001": 1, "LBL001": 1, "I18N001": 1, "UX001": 1, "UX002": 1}
    )


def test_larger_mixed_retains_native_and_one_first_fatal_pyxform_result(tmp_path, monkeypatch) -> None:
    import xlsform_lint.pyxform_adapter as adapter

    path = _write_large(tmp_path / "larger-mixed.xlsx", lint=True, fatal=True)
    original = adapter.convert_workbook
    calls = 0

    def counted(data: bytes):
        nonlocal calls
        calls += 1
        return original(data)

    monkeypatch.setattr(adapter, "convert_workbook", counted)
    diagnostics = lint_workbook(path)
    assert calls == 1
    assert [item.rule_id for item in diagnostics] == [
        "LBL001", "I18N001", "UX001", "UX002", "PX002", "CHO001"
    ]
    assert sum(item.rule_id.startswith("PX") for item in diagnostics) == 1
    assert main([str(path)]) == 1
