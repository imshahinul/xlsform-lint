from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from xlsform_lint.diagnostics import DiagnosticOrigin, Severity, sort_diagnostics
from xlsform_lint.engine import lint_workbook
from xlsform_lint.native_model import (
    SelectKind,
    SurveyTypeCategory,
    build_native_model,
    classify_select,
    classify_survey_type,
)
from xlsform_lint.native_rules import NATIVE_RULES, run_native_rules


def _rules(path):
    return tuple(item.rule_id for item in lint_workbook(path))


def _save_bytes(survey_rows, choices_rows=None, survey_title="Survey", choices_title="Choices"):
    workbook = Workbook()
    survey = workbook.active
    survey.title = survey_title
    for row in survey_rows:
        survey.append(row)
    if choices_rows is not None:
        choices = workbook.create_sheet(choices_title)
        for row in choices_rows:
            choices.append(row)
    settings = workbook.create_sheet("Settings")
    settings.append(("form_title", "form_id", "version"))
    settings.append(("Phase 3", "phase_3", "1"))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_native_model_preserves_physical_identity_and_classifies_conservatively():
    data = _save_bytes(
        [("type", "name", "label::English", "label::French"), ("text", "q", "Q", None)],
        [("list_name", "name", "label"), ("yesno", "yes", "Yes")],
        survey_title="SuRvEy", choices_title="ChOiCeS",
    )
    model = build_native_model(data)
    assert model.survey_sheet.physical_name == "SuRvEy"
    assert model.choices_sheet.physical_name == "ChOiCeS"
    assert [(column.language, column.actual) for column in model.survey_sheet.translations] == [
        ("English", "label::English"), ("French", "label::French")
    ]
    assert model.survey_rows[0].row == 2
    assert model.survey_rows[0].category is SurveyTypeCategory.VISIBLE
    assert classify_survey_type("mystery") is SurveyTypeCategory.UNKNOWN
    assert classify_select("select_one yesno") == (SelectKind.INTERNAL_STATIC, "yesno")
    assert classify_select("select_one_from_file cities.csv")[0] is SelectKind.EXTERNAL_FILE


def test_cho001_matrix(tmp_path, xlsform_factory):
    header = [("type", "name", "label")]
    choices = [("list_name", "name", "label"), ("used", "a", "A"), ("unused", "b", "B")]
    unused = xlsform_factory(tmp_path / "unused.xlsx", header + [("text", "q", "Q")], choices_rows=choices)
    assert _rules(unused).count("CHO001") == 2
    one = xlsform_factory(tmp_path / "one.xlsx", header + [("select_one used", "q", "Q")], choices_rows=choices)
    assert _rules(one).count("CHO001") == 1
    multiple = xlsform_factory(tmp_path / "multiple.xlsx", header + [("select_multiple used", "q", "Q")], choices_rows=choices)
    assert _rules(multiple).count("CHO001") == 1
    for index, select_type in enumerate(("select_one_from_file cities.csv", "select_one ${list}", "select_weird")):
        path = xlsform_factory(tmp_path / f"safe-{index}.xlsx", header + [(select_type, "q", "Q")], choices_rows=choices)
        assert "CHO001" not in _rules(path)


def test_lbl001_type_and_translation_matrix(tmp_path, xlsform_factory):
    rows = [
        ("type", "name", "label"), ("text", "missing", None), ("text", "ok", "OK"),
        ("note", "note_missing", None), ("calculate", "calc", None), ("hidden", "hide", None),
        ("start", "start", None), ("begin group", "g", None), ("mystery", "u", None),
    ]
    path = xlsform_factory(tmp_path / "labels.xlsx", rows)
    diagnostics = [item for item in lint_workbook(path) if item.rule_id == "LBL001"]
    assert [item.source.row for item in diagnostics] == [2, 4]
    multilingual = xlsform_factory(
        tmp_path / "multilingual.xlsx",
        [("type", "name", "label::English", "label::French"),
         ("text", "partial", "Hello", None), ("text", "empty", None, " ")],
    )
    assert _rules(multilingual).count("I18N001") == 1
    assert _rules(multilingual).count("LBL001") == 1


def test_lbl002_and_choice_i18n_noise_matrix(tmp_path, xlsform_factory):
    path = xlsform_factory(
        tmp_path / "choice-labels.xlsx", [("type", "name", "label"), ("text", "q", "Q")],
        choices_rows=[("list_name", "name", "label"), ("region", "north", None),
                      ("region", "south", "South"), (None, None, None), ("region", None, None)],
    )
    assert _rules(path).count("LBL002") == 1
    multilingual = xlsform_factory(
        tmp_path / "choice-i18n.xlsx", [("type", "name", "label"), ("select_one region", "q", "Q")],
        choices_rows=[("list_name", "name", "label::English", "label::French"),
                      ("region", "north", "North", None), ("region", "empty", None, None)],
    )
    assert _rules(multilingual).count("I18N001") == 1
    assert _rules(multilingual).count("LBL002") == 1


def test_i18n001_survey_matrix(tmp_path, xlsform_factory):
    path = xlsform_factory(
        tmp_path / "i18n.xlsx",
        [("type", "name", "label::English", "label::French", "foo::German"),
         ("text", "fr_missing", "English", None, None),
         ("text", "en_missing", None, "Français", None),
         ("text", "complete", "Both", "Deux", None),
         ("text", "blank", None, None, "not a label"),
         ("calculate", "calc", "Calculation", None, None)],
    )
    findings = [item for item in lint_workbook(path) if item.rule_id == "I18N001"]
    assert [(item.source.row, item.source.column) for item in findings] == [
        (2, "label::French"), (3, "label::English")
    ]
    single = xlsform_factory(tmp_path / "single.xlsx", [("type", "name", "label::English"), ("text", "q", None)])
    assert "I18N001" not in _rules(single)


def test_translation_order_and_plain_plus_translated_policy(tmp_path, xlsform_factory):
    path = xlsform_factory(
        tmp_path / "translation-order.xlsx",
        [("type", "name", "label", "label::French", "label::English", "label::Spanish"),
         ("text", "q", "Plain fallback", "Français", None, None)],
    )
    findings = [item for item in lint_workbook(path) if item.rule_id == "I18N001"]
    assert [item.source.column for item in findings] == ["label::English", "label::Spanish"]
    assert "LBL001" not in _rules(path)


def test_ux001_matrix(tmp_path, xlsform_factory):
    path = xlsform_factory(
        tmp_path / "constraints.xlsx",
        [("type", "name", "label", "constraint", "constraint_message", "constraint_message::French"),
         ("integer", "missing", "Missing", ". > 0", None, None),
         ("integer", "plain", "Plain", ". > 0", "Positive", None),
         ("integer", "translated", "Translated", ". > 0", None, "Positif"),
         ("integer", "none", "None", " ", None, None)],
    )
    findings = [item for item in lint_workbook(path) if item.rule_id == "UX001"]
    assert [item.source.row for item in findings] == [2]


def test_ord001_is_choice_filter_only_and_unambiguous(tmp_path, xlsform_factory):
    header = ("type", "name", "label", "choice_filter", "relevant", "constraint", "calculation")
    path = xlsform_factory(
        tmp_path / "order.xlsx",
        [header, ("text", "early", "Early", None, None, None, None),
         ("select_one list", "current", "Current", "x=${later} and y=${early}", None, None, None),
         ("text", "other", "Other", None, "${later}", "${later} > 0", "${later}"),
         ("text", "later", "Later", None, None, None, None)],
        choices_rows=[("list_name", "name", "label"), ("list", "a", "A")],
    )
    findings = [item for item in lint_workbook(path) if item.rule_id == "ORD001"]
    assert len(findings) == 1 and "${later}" in findings[0].message
    duplicate = xlsform_factory(
        tmp_path / "ambiguous.xlsx", [header, ("text", "q", "Q", "x=${dup}", None, None, None),
                                      ("text", "dup", "One", None, None, None, None),
                                      ("text", "dup", "Two", None, None, None, None)],
    )
    assert "ORD001" not in _rules(duplicate)

    grouped = xlsform_factory(
        tmp_path / "grouped.xlsx",
        [("type", "name", "label", "choice_filter"),
         ("begin group", "group", "Group", None),
         ("select_one list", "q", "Q", "x=${later}"),
         ("end group", None, None, None),
         ("text", "later", "Later", None)],
        choices_rows=[("list_name", "name", "label"), ("list", "a", "A")],
    )
    assert "ORD001" not in _rules(grouped)


def test_ux002_matrix_and_severity(tmp_path, xlsform_factory):
    path = xlsform_factory(
        tmp_path / "defaults.xlsx",
        [("type", "name", "label", "required", "default"),
         ("text", "positive", "Positive", " YES ", "value"),
         ("text", "blank_default", "Blank", "yes", " "),
         ("text", "blank_required", "Blank required", None, "value"),
         ("text", "no", "No", "no", "value")],
    )
    findings = [item for item in lint_workbook(path) if item.rule_id == "UX002"]
    assert len(findings) == 1 and findings[0].severity is Severity.INFO


def test_duplicate_headers_skip_dependent_rules():
    model = build_native_model(_save_bytes([("type", "name", "label", "label"), ("text", "q", None, None)]))
    assert run_native_rules(model, "ambiguous.xlsx") == ()


def test_native_executor_has_exactly_seven_rules_and_global_sort_is_order_independent():
    assert len(NATIVE_RULES) == 7
    data = _save_bytes(
        [("type", "name", "label", "constraint", "constraint_message", "required", "default"),
         ("text", "q", None, ". != ''", None, "yes", "x")],
        [("list_name", "name", "label"), ("unused", "a", None)],
    )
    model = build_native_model(data)
    diagnostics = run_native_rules(model, "mixed.xlsx")
    assert {item.rule_id for item in diagnostics} == {"CHO001", "LBL001", "LBL002", "UX001", "UX002"}
    assert all(item.origin is DiagnosticOrigin.NATIVE for item in diagnostics)
    assert sort_diagnostics(diagnostics) == sort_diagnostics(reversed(diagnostics))


def test_native_diagnostic_survives_independent_pyxform_failure(tmp_path, xlsform_factory):
    path = xlsform_factory(
        tmp_path / "native-plus-px.xlsx",
        [("type", "name", "label", "constraint", "constraint_message"),
         ("text", "good", "Good", ". != ''", None), ("text", "1bad", "Bad", None, None)],
    )
    diagnostics = lint_workbook(path)
    assert [item.rule_id for item in diagnostics] == ["UX001", "PX002"]
    assert diagnostics[0].origin is DiagnosticOrigin.NATIVE
    assert diagnostics[1].origin is DiagnosticOrigin.PYXFORM
