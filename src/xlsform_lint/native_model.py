"""Bounded read-only workbook facts for the seven Phase 3 native rules."""

from __future__ import annotations

from dataclasses import dataclass
from enum import Enum
from io import BytesIO
import re
from types import MappingProxyType
from typing import Any, Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


_TRANSLATION_HEADER = re.compile(r"^label::(?P<language>.+)$", re.IGNORECASE)
_STATIC_SELECT = re.compile(
    r"^select_(?:one|multiple)\s+(?P<list>[A-Za-z0-9_.-]+)$", re.IGNORECASE
)


class SurveyTypeCategory(str, Enum):
    VISIBLE = "ordinary-visible-question"
    NOTE = "note"
    CALCULATE = "calculate"
    HIDDEN = "hidden"
    METADATA = "metadata-internal"
    GROUP_REPEAT = "group-repeat"
    UNKNOWN = "unknown"


class SelectKind(str, Enum):
    INTERNAL_STATIC = "internal-static"
    INTERNAL_DYNAMIC = "internal-dynamic"
    EXTERNAL_FILE = "external-file"
    UNKNOWN = "unknown"


@dataclass(frozen=True)
class HeaderCell:
    actual: str
    logical: str
    index: int
    cell: str


@dataclass(frozen=True)
class TranslationColumn:
    base: str
    language: str
    actual: str
    index: int


@dataclass(frozen=True)
class SheetFacts:
    canonical_name: str
    physical_name: str
    headers: tuple[HeaderCell, ...]
    unique_headers: Mapping[str, HeaderCell]
    ambiguous_headers: frozenset[str]
    translations: tuple[TranslationColumn, ...]


@dataclass(frozen=True)
class SurveyRow:
    row: int
    values: tuple[Any, ...]
    sheet: SheetFacts
    category: SurveyTypeCategory


@dataclass(frozen=True)
class ChoiceRow:
    row: int
    values: tuple[Any, ...]
    sheet: SheetFacts


@dataclass(frozen=True)
class SelectUsage:
    row: int
    kind: SelectKind
    list_name: str | None


@dataclass(frozen=True)
class WorkbookFacts:
    survey_sheet: SheetFacts | None
    choices_sheet: SheetFacts | None
    survey_rows: tuple[SurveyRow, ...]
    choice_rows: tuple[ChoiceRow, ...]
    survey_by_name: Mapping[str, tuple[SurveyRow, ...]]
    choice_rows_by_list: Mapping[str, tuple[ChoiceRow, ...]]
    select_usage: tuple[SelectUsage, ...]


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def semantic_text(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    value = value.strip()
    return value or None


def classify_survey_type(value: Any) -> SurveyTypeCategory:
    text = semantic_text(value)
    if text is None:
        return SurveyTypeCategory.UNKNOWN
    lowered = text.casefold()
    if lowered == "note":
        return SurveyTypeCategory.NOTE
    if lowered == "calculate":
        return SurveyTypeCategory.CALCULATE
    if lowered == "hidden":
        return SurveyTypeCategory.HIDDEN
    if lowered in {"start", "end", "today", "deviceid", "username", "email", "audit", "background-audio"}:
        return SurveyTypeCategory.METADATA
    if lowered in {"begin group", "end group", "begin repeat", "end repeat"}:
        return SurveyTypeCategory.GROUP_REPEAT
    visible_exact = {
        "text", "integer", "decimal", "date", "time", "datetime", "geopoint",
        "geotrace", "geoshape", "image", "audio", "video", "file", "barcode",
        "range", "rank", "acknowledge", "xml-external",
    }
    if lowered in visible_exact or _STATIC_SELECT.fullmatch(text):
        return SurveyTypeCategory.VISIBLE
    return SurveyTypeCategory.UNKNOWN


def classify_select(value: Any) -> tuple[SelectKind, str | None] | None:
    text = semantic_text(value)
    if text is None or not text.casefold().startswith("select_"):
        return None
    static = _STATIC_SELECT.fullmatch(text)
    if static:
        return SelectKind.INTERNAL_STATIC, static["list"]
    lowered = text.casefold()
    if "external" in lowered or " from file" in lowered or lowered.endswith(".csv"):
        return SelectKind.EXTERNAL_FILE, None
    if "(" in text or "${" in text or " " in text:
        return SelectKind.INTERNAL_DYNAMIC, None
    return SelectKind.UNKNOWN, None


def row_value(row: SurveyRow | ChoiceRow, logical_header: str) -> Any:
    header = row.sheet.unique_headers.get(logical_header.casefold())
    if header is None or header.index >= len(row.values):
        return None
    return row.values[header.index]


def _sheet_facts(worksheet, canonical_name: str) -> SheetFacts:
    raw_headers = tuple(cell.value for cell in next(worksheet.iter_rows(max_row=1), ()))
    headers = tuple(
        HeaderCell(value, value.strip().casefold(), index, f"{get_column_letter(index + 1)}1")
        for index, value in enumerate(raw_headers)
        if isinstance(value, str) and value.strip()
    )
    counts = {header.logical: sum(item.logical == header.logical for item in headers) for header in headers}
    unique = MappingProxyType(
        {header.logical: header for header in headers if counts[header.logical] == 1}
    )
    translations = []
    for header in headers:
        match = _TRANSLATION_HEADER.fullmatch(header.actual.strip())
        if match and match["language"].strip():
            translations.append(
                TranslationColumn("label", match["language"], header.actual, header.index)
            )
    return SheetFacts(
        canonical_name=canonical_name,
        physical_name=worksheet.title,
        headers=headers,
        unique_headers=unique,
        ambiguous_headers=frozenset(name for name, count in counts.items() if count > 1),
        translations=tuple(translations),
    )


def build_native_model(workbook_bytes: bytes) -> WorkbookFacts:
    """Build physical facts and only the indexes required by frozen native rules."""

    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    try:
        canonical = {}
        for worksheet in workbook.worksheets:
            name = worksheet.title.casefold()
            if name in {"survey", "choices", "settings"} and name not in canonical:
                canonical[name] = worksheet

        survey_sheet = _sheet_facts(canonical["survey"], "survey") if "survey" in canonical else None
        choices_sheet = _sheet_facts(canonical["choices"], "choices") if "choices" in canonical else None

        survey_rows: list[SurveyRow] = []
        if survey_sheet is not None:
            worksheet = canonical["survey"]
            for physical_row, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
                values = tuple(cells)
                temporary = SurveyRow(physical_row, values, survey_sheet, SurveyTypeCategory.UNKNOWN)
                survey_rows.append(
                    SurveyRow(physical_row, values, survey_sheet, classify_survey_type(row_value(temporary, "type")))
                )

        choice_rows: list[ChoiceRow] = []
        if choices_sheet is not None:
            worksheet = canonical["choices"]
            choice_rows = [
                ChoiceRow(physical_row, tuple(cells), choices_sheet)
                for physical_row, cells in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2)
            ]

        survey_by_name_mut: dict[str, list[SurveyRow]] = {}
        for row in survey_rows:
            name = semantic_text(row_value(row, "name"))
            if name is not None:
                survey_by_name_mut.setdefault(name, []).append(row)
        choice_by_list_mut: dict[str, list[ChoiceRow]] = {}
        for row in choice_rows:
            list_name = semantic_text(row_value(row, "list_name"))
            if list_name is not None:
                choice_by_list_mut.setdefault(list_name, []).append(row)
        usages = []
        for row in survey_rows:
            classified = classify_select(row_value(row, "type"))
            if classified is not None:
                usages.append(SelectUsage(row.row, *classified))

        return WorkbookFacts(
            survey_sheet,
            choices_sheet,
            tuple(survey_rows),
            tuple(choice_rows),
            MappingProxyType({key: tuple(value) for key, value in sorted(survey_by_name_mut.items())}),
            MappingProxyType({key: tuple(value) for key, value in sorted(choice_by_list_mut.items())}),
            tuple(usages),
        )
    finally:
        workbook.close()
