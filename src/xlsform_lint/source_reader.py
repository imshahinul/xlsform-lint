"""Read-only source indexing for deterministic workbook enrichment."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class SourceCell:
    """A non-empty physical cell under a uniquely named logical header."""

    sheet: str
    row: int
    column: str
    cell: str
    value: Any


@dataclass(frozen=True)
class WorkbookSourceIndex:
    """Minimum source facts used only after pyxform establishes invalidity."""

    cells: tuple[SourceCell, ...]

    def find(
        self,
        *,
        sheet: str,
        row: int | None = None,
        column: str | None = None,
    ) -> tuple[SourceCell, ...]:
        return tuple(
            cell
            for cell in self.cells
            if cell.sheet.casefold() == sheet.casefold()
            and (row is None or cell.row == row)
            and (column is None or cell.column == column)
        )


def read_source_index(workbook_bytes: bytes) -> WorkbookSourceIndex:
    """Index physical cells from an independent in-memory workbook copy.

    Duplicate or blank headers are deliberately not assigned a logical column.
    This prevents source enrichment from claiming precision it cannot prove.
    """

    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    cells: list[SourceCell] = []
    try:
        for worksheet in workbook.worksheets:
            headers = tuple(
                cell.value for cell in next(worksheet.iter_rows(max_row=1), ())
            )
            normalized = tuple(
                value.strip().casefold() if isinstance(value, str) else None
                for value in headers
            )
            counts = {name: normalized.count(name) for name in normalized if name}
            for row in worksheet.iter_rows(min_row=2):
                for index, cell in enumerate(row):
                    header = normalized[index] if index < len(normalized) else None
                    if header and counts[header] == 1 and cell.value is not None:
                        cells.append(
                            SourceCell(
                                sheet=worksheet.title,
                                row=cell.row,
                                column=header,
                                cell=f"{get_column_letter(cell.column)}{cell.row}",
                                value=cell.value,
                            )
                        )
        return WorkbookSourceIndex(tuple(cells))
    finally:
        workbook.close()


def read_source_cell(
    workbook_bytes: bytes, sheet_name: str, row: int, column: int
) -> Any:
    """Read one physical cell from an independent in-memory workbook copy."""

    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    try:
        return workbook[sheet_name].cell(row=row, column=column).value
    finally:
        workbook.close()


def read_sheet_names(workbook_bytes: bytes) -> tuple[str, ...]:
    """Return physical sheet names from an independent in-memory copy."""

    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()
